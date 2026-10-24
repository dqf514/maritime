from datetime import datetime, timezone
from uuid import UUID, uuid4
import hashlib
import secrets

from fastapi import APIRouter, Depends, File, HTTPException, Request, Response, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_db
from app.models import (
    ApiKey,
    BackupJob,
    MigrationJob,
    MigrationProposal,
    MigrationSource,
    Module,
    SelfCheckResult,
    SelfCheckRun,
    Tenant,
    TenantModuleLicense,
    User,
)
from app.schemas import (
    ApiKeyCreate,
    ApiKeyOut,
    BackupJobOut,
    LicenseOut,
    LoginIn,
    MeOut,
    MigrationJobCreate,
    MigrationJobOut,
    MigrationProposalOut,
    MigrationSourceIn,
    SearchHit,
    SelfCheckRunOut,
    SelfCheckResultOut,
    TenantOut,
    TokenOut,
    UserOut,
)
from app.security import (
    AuthContext,
    create_access_token,
    get_current_auth,
    require_module,
    set_session_cookie,
    verify_password,
)
from app.services.audit import audit
from app.services.ratelimit import rate_limit
from app.services.recycle import soft_delete
from app.services.search_acl import allowed_omni_hits, filter_omni_query
from checks.registry import run_all_checks, score_results

router = APIRouter()


@router.post("/auth/login", response_model=TokenOut, tags=["Auth"])
def login(
    body: LoginIn,
    response: Response,
    request: Request,
    db: Session = Depends(get_db),
    _rl: None = Depends(rate_limit("auth.login", max_hits=5, window_seconds=60)),
):
    from app.services.identity import ensure_tenant_policy, now_utc_naive

    ip = request.client.host if request.client else None
    tenant = db.scalar(select(Tenant).where(Tenant.code == body.tenant_code))
    if not tenant:
        # commit=True: the 401 below would otherwise roll the audit row back
        audit(
            db,
            action="auth.login_failed",
            detail={"email": body.email, "tenant_code": body.tenant_code, "reason": "unknown_tenant"},
            ip=ip,
            commit=True,
        )
        raise HTTPException(status_code=401, detail="Invalid credentials")
    if tenant.status == "suspended":
        raise HTTPException(status_code=403, detail={"code": "TENANT_SUSPENDED", "message": "Tenant suspended"})
    policy = ensure_tenant_policy(db, tenant.id)
    if not policy.password_enabled:
        raise HTTPException(status_code=403, detail={"code": "PASSWORD_DISABLED", "message": "Password login disabled"})
    user = db.scalar(select(User).where(User.tenant_id == tenant.id, User.email == body.email))
    if not user or not user.password_hash or not verify_password(body.password, user.password_hash):
        audit(
            db,
            tenant_id=tenant.id,
            actor_user_id=user.id if user else None,
            action="auth.login_failed",
            entity_type="user",
            entity_id=user.id if user else None,
            detail={"email": body.email, "reason": "bad_credentials"},
            ip=ip,
            commit=True,
        )
        raise HTTPException(status_code=401, detail="Invalid credentials")
    if user.status != "active":
        audit(
            db,
            tenant_id=tenant.id,
            actor_user_id=user.id,
            action="auth.login_failed",
            entity_type="user",
            entity_id=user.id,
            detail={"email": body.email, "reason": "user_inactive"},
            ip=ip,
            commit=True,
        )
        raise HTTPException(status_code=403, detail="User inactive")
    if policy.require_email_verify and not user.email_verified_at:
        raise HTTPException(
            status_code=403,
            detail={"code": "EMAIL_NOT_VERIFIED", "message": "Email verification required before sign-in"},
        )
    user.last_login_at = now_utc_naive()
    # Audit row rides this commit, atomically with last_login_at
    audit(
        db,
        tenant_id=tenant.id,
        actor_user_id=user.id,
        action="auth.login_success",
        entity_type="user",
        entity_id=user.id,
        detail={"email": user.email},
        ip=ip,
    )
    db.commit()
    token = create_access_token(user_id=user.id, tenant_id=tenant.id, email=user.email, pwv=user.password_version or 1)
    # Browser session: token is also planted as an HttpOnly cookie; the JSON
    # body keeps access_token for backward compatibility with header clients.
    set_session_cookie(response, token)
    return TokenOut(access_token=token)


@router.get("/me", response_model=MeOut, tags=["Auth"])
def me(auth: AuthContext = Depends(get_current_auth), db: Session = Depends(get_db)):
    tenant = db.get(Tenant, auth.tenant_id)
    assert tenant
    licenses = db.scalars(
        select(TenantModuleLicense.module_code).where(
            TenantModuleLicense.tenant_id == auth.tenant_id,
            TenantModuleLicense.status.in_(["active", "grace"]),
        )
    ).all()
    return MeOut(
        user=UserOut(
            id=auth.user.id,
            email=auth.user.email,
            full_name=auth.user.full_name,
            locale=auth.user.locale,
            roles=auth.roles,
            email_verified=bool(auth.user.email_verified_at),
        ),
        tenant=TenantOut(
            id=tenant.id,
            name=tenant.name,
            code=tenant.code,
            status=tenant.status,
            default_locale=tenant.default_locale,
            default_timezone=tenant.default_timezone,
            profile_tier=tenant.profile_tier,
        ),
        licensed_modules=list(licenses),
    )


@router.get("/tenants/current", response_model=TenantOut, tags=["Tenants"])
def current_tenant(auth: AuthContext = Depends(get_current_auth), db: Session = Depends(get_db)):
    tenant = db.get(Tenant, auth.tenant_id)
    assert tenant
    return TenantOut(
        id=tenant.id,
        name=tenant.name,
        code=tenant.code,
        status=tenant.status,
        default_locale=tenant.default_locale,
        default_timezone=tenant.default_timezone,
        profile_tier=tenant.profile_tier,
    )


@router.get("/tenants/current/licenses", response_model=list[LicenseOut], tags=["Licenses"])
def list_licenses(auth: AuthContext = Depends(get_current_auth), db: Session = Depends(get_db)):
    rows = db.execute(
        select(TenantModuleLicense, Module)
        .join(Module, Module.code == TenantModuleLicense.module_code)
        .where(TenantModuleLicense.tenant_id == auth.tenant_id)
    ).all()
    return [
        LicenseOut(module_code=lic.module_code, status=lic.status, is_core=mod.is_core)
        for lic, mod in rows
    ]


@router.get("/search", response_model=list[SearchHit], tags=["Shell"])
def omni_search(
    q: str = "",
    auth: AuthContext = Depends(get_current_auth),
    db: Session = Depends(get_db),
):
    """Return only pages/commands the current role (and tenant licenses) may open.

    U/闭环：q 非空时追加实体记录命中（对手方/联系人/船舶/港口），直达详情。
    """
    allowed = allowed_omni_hits(db, tenant_id=auth.tenant_id, roles=auth.roles)
    hits = filter_omni_query(allowed, q)
    if q.strip():
        hits = hits + _entity_hits(db, auth, q.strip())
    return hits


def _entity_hits(db: Session, auth: AuthContext, q: str, limit: int = 5) -> list:
    """实体记录命中：对手方/联系人/船舶/港口（按租户过滤）。"""
    from sqlalchemy import or_

    from app.models_wave1 import Counterparty, CounterpartyContact, Port, Vessel
    from app.models import User  # noqa: F401 — 保持导入完整

    like = f"%{q}%"
    out: list = []

    def _hit(id_: str, title: str, module: str, href: str, kw: list[str]):
        out.append(SearchHit(id=id_, title=title, keywords=kw, module=module, href=href))

    try:
        for r in db.scalars(
            select(Counterparty)
            .where(Counterparty.tenant_id == auth.tenant_id, Counterparty.deleted_at.is_(None), Counterparty.name.ilike(like))
            .limit(limit)
        ).all():
            _hit(f"cp-{r.id}", r.name, "masterdata", f"/masterdata/counterparties?party={r.id}", ["counterparty", "对手方"])
        for r in db.execute(
            select(CounterpartyContact, Counterparty.name)
            .join(Counterparty, CounterpartyContact.counterparty_id == Counterparty.id)
            .where(
                CounterpartyContact.tenant_id == auth.tenant_id,
                or_(CounterpartyContact.name.ilike(like), CounterpartyContact.email.ilike(like)),
            )
            .limit(limit)
        ).all():
            contact, party_name = r
            _hit(f"ct-{contact.id}", f"{contact.name} @ {party_name}", "masterdata", f"/masterdata/counterparties?party={contact.counterparty_id}", ["contact", "联系人"])
        for r in db.scalars(
            select(Vessel)
            .where(Vessel.tenant_id == auth.tenant_id, Vessel.deleted_at.is_(None), Vessel.name.ilike(like))
            .limit(limit)
        ).all():
            _hit(f"vs-{r.id}", r.name, "masterdata", "/masterdata/vessels", ["vessel", "船舶"])
        for r in db.scalars(
            select(Port)
            .where(Port.deleted_at.is_(None), or_(Port.name.ilike(like), Port.unlocode.ilike(like)))
            .limit(limit)
        ).all():
            _hit(f"pt-{r.id}", f"{r.name} ({r.unlocode})", "masterdata", "/masterdata/ports", ["port", "港口"])
    except Exception:  # noqa: BLE001 — 搜索失败不影响页面命中
        pass
    return out


@router.post("/settings/selfcheck/run", response_model=SelfCheckRunOut, tags=["SelfCheck"])
def run_selfcheck(
    auth: AuthContext = Depends(require_module("selfcheck")),
    db: Session = Depends(get_db),
):
    run = SelfCheckRun(tenant_id=auth.tenant_id, triggered_by=auth.user_id, status="running")
    db.add(run)
    db.flush()
    results = run_all_checks(db)
    for r in results:
        db.add(
            SelfCheckResult(
                run_id=run.id,
                check_id=r.check_id,
                severity=r.severity,
                status=r.status,
                message=r.message,
                details=r.details or {},
            )
        )
    run.score = score_results(results)
    run.status = "completed"
    run.finished_at = datetime.now(timezone.utc)
    db.commit()
    db.refresh(run)
    return SelfCheckRunOut(
        id=run.id,
        status=run.status,
        score=run.score,
        started_at=run.started_at,
        finished_at=run.finished_at,
        results=[
            SelfCheckResultOut(
                check_id=r.check_id, severity=r.severity, status=r.status, message=r.message
            )
            for r in results
        ],
    )


@router.post("/settings/dataops/backups", response_model=BackupJobOut, tags=["DataOps"])
def create_backup(
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    """全量租户数据备份（所有业务表行级 JSON dump）。

    Wave-0 的配置摘要备份已由全量 dump 取代；文件仍写
    ``data/backups/<tenant>/<job>.json``（不在 /uploads 静态挂载下）。
    """
    from app.services.tenant_backup import dump_tenant, write_backup_file

    payload = dump_tenant(db, auth.tenant_id, created_by=auth.user_id, trigger="manual")
    job = write_backup_file(db, auth.tenant_id, payload, created_by=auth.user_id)
    return BackupJobOut(
        id=job.id,
        status=job.status,
        trigger=job.trigger,
        storage_path=job.storage_path,
        checksum=job.checksum,
        created_at=job.created_at,
        finished_at=job.finished_at,
        error=job.error,
    )


@router.get("/settings/dataops/backups", response_model=list[BackupJobOut], tags=["DataOps"])
def list_backups(
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    rows = db.scalars(
        select(BackupJob)
        .where(BackupJob.tenant_id == auth.tenant_id)
        .order_by(BackupJob.created_at.desc())
        .limit(50)
    ).all()
    return [
        BackupJobOut(
            id=r.id,
            status=r.status,
            trigger=r.trigger,
            storage_path=r.storage_path,
            checksum=r.checksum,
            created_at=r.created_at,
            finished_at=r.finished_at,
            error=r.error,
        )
        for r in rows
    ]


class BackupJobPatchIn(BaseModel):
    status: str | None = None
    error: str | None = None


@router.patch("/settings/dataops/backups/{job_id}", response_model=BackupJobOut, tags=["DataOps"])
def patch_backup(
    job_id: UUID,
    body: BackupJobPatchIn,
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    """Annotate a backup job record (status / error note). BackupJob has no enabled/schedule fields."""
    row = db.get(BackupJob, job_id)
    if not row or row.tenant_id != auth.tenant_id:
        raise HTTPException(404, "Backup not found")
    if body.status is not None and body.status != "deleted":
        row.status = body.status
    if body.error is not None:
        row.error = body.error
    db.commit()
    db.refresh(row)
    return BackupJobOut(
        id=row.id,
        status=row.status,
        trigger=row.trigger,
        storage_path=row.storage_path,
        checksum=row.checksum,
        created_at=row.created_at,
        finished_at=row.finished_at,
        error=row.error,
    )


@router.delete("/settings/dataops/backups/{job_id}", tags=["DataOps"])
def delete_backup(
    job_id: UUID,
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    """Delete a backup job and its file — same purge path as backup_scheduler prune."""
    row = db.get(BackupJob, job_id)
    if not row or row.tenant_id != auth.tenant_id:
        raise HTTPException(404, "Backup not found")
    if row.status in {"queued", "running"}:
        raise HTTPException(409, detail={"code": "IN_USE", "message": "Backup is still in progress"})
    if row.storage_path:
        from pathlib import Path

        from app.services.tenant_backup import _data_root

        try:
            p = _data_root() / row.storage_path
            if p.is_file():
                p.unlink()
        except OSError:  # pragma: no cover - 文件已不存在等
            pass
    db.delete(row)
    db.commit()
    return {"ok": True}


# ── 自动备份设置（每租户一行；由 backup_scheduler 周期 tick 执行） ──


def _schedule_out(s) -> dict:
    return {
        "enabled": s.enabled,
        "frequency": s.frequency,
        "hour_utc": s.hour_utc,
        "weekday": s.weekday,
        "keep_last": s.keep_last,
        "last_run_at": s.last_run_at.isoformat() if s.last_run_at else None,
        "next_run_at": s.next_run_at.isoformat() if s.next_run_at else None,
    }


@router.get("/settings/dataops/backup-schedule", tags=["DataOps"])
def get_backup_schedule(
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    from app.models import BackupSchedule

    row = db.scalar(select(BackupSchedule).where(BackupSchedule.tenant_id == auth.tenant_id))
    if not row:
        # 未配置时返回默认值（disabled），前端据此渲染表单
        return {
            "enabled": False,
            "frequency": "daily",
            "hour_utc": 2,
            "weekday": 0,
            "keep_last": 14,
            "last_run_at": None,
            "next_run_at": None,
        }
    return _schedule_out(row)


class BackupScheduleIn(BaseModel):
    enabled: bool
    frequency: str  # daily | weekly
    hour_utc: int = Field(ge=0, le=23)
    weekday: int = Field(ge=0, le=6, default=0)
    keep_last: int = Field(ge=1, le=365, default=14)


@router.put("/settings/dataops/backup-schedule", tags=["DataOps"])
def put_backup_schedule(
    body: BackupScheduleIn,
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    from app.models import BackupSchedule
    from app.services.backup_scheduler import compute_next_run

    if body.frequency not in ("daily", "weekly"):
        raise HTTPException(422, detail={"code": "INVALID_FREQUENCY"})
    row = db.scalar(select(BackupSchedule).where(BackupSchedule.tenant_id == auth.tenant_id))
    if not row:
        row = BackupSchedule(tenant_id=auth.tenant_id)
        db.add(row)
    row.enabled = body.enabled
    row.frequency = body.frequency
    row.hour_utc = body.hour_utc
    row.weekday = body.weekday
    row.keep_last = body.keep_last
    row.next_run_at = compute_next_run(row, datetime.now(timezone.utc)) if body.enabled else None
    db.commit()
    db.refresh(row)
    return _schedule_out(row)


@router.post("/settings/dataops/migrations", response_model=MigrationJobOut, tags=["DataOps"])
def create_migration(
    body: MigrationJobCreate,
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    job = MigrationJob(
        tenant_id=auth.tenant_id,
        status="draft",
        summary={"note": body.note},
        created_by=auth.user_id,
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    return MigrationJobOut(id=job.id, status=job.status, summary=job.summary, created_at=job.created_at)


@router.post("/settings/dataops/migrations/{job_id}/sources", tags=["DataOps"])
def add_migration_source(
    job_id: UUID,
    body: MigrationSourceIn,
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    job = db.get(MigrationJob, job_id)
    if not job or job.tenant_id != auth.tenant_id:
        raise HTTPException(status_code=404, detail="Migration job not found")
    src = MigrationSource(job_id=job.id, source_type=body.source_type, config=body.config, status="ready")
    db.add(src)
    db.commit()
    return {"id": str(src.id), "source_type": src.source_type, "status": src.status}


@router.post("/settings/dataops/migrations/{job_id}/run-analyze", tags=["DataOps"])
def run_analyze(
    job_id: UUID,
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    job = db.get(MigrationJob, job_id)
    if not job or job.tenant_id != auth.tenant_id:
        raise HTTPException(status_code=404, detail="Migration job not found")
    # Clear prior proposed rows for re-run
    for old in db.scalars(select(MigrationProposal).where(MigrationProposal.job_id == job.id)).all():
        if old.status == "proposed":
            db.delete(old)
    samples = [
        ("vessel", "create", 0.92, {"name": "MV OCEAN STAR", "imo": "9123456", "vessel_type": "Bulker"}),
        ("counterparty", "create", 0.88, {"name": "Pacific Charterers Ltd", "type": "charterer", "country": "SG"}),
        ("port", "create", 0.95, {"name": "Singapore", "unlocode": "SGSIN", "timezone": "Asia/Singapore"}),
        ("voyage", "create", 0.71, {"voyage_no": "V2026-001", "cargo": "Coal"}),
    ]
    for entity_type, action, conf, payload in samples:
        db.add(
            MigrationProposal(
                job_id=job.id,
                entity_type=entity_type,
                action=action,
                confidence=conf,
                payload=payload,
                status="proposed",
            )
        )
    job.status = "review"
    job.summary = {**job.summary, "proposal_count": len(samples)}
    db.commit()
    return {"status": job.status, "proposal_count": len(samples)}


@router.post("/settings/dataops/migrations/{job_id}/upload-excel", tags=["DataOps"])
async def upload_excel(
    job_id: UUID,
    file: UploadFile = File(...),
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    from openpyxl import load_workbook
    import io

    from app.services.upload_policy import get_upload_policy

    policy = get_upload_policy(db, "document")
    MAX_EXCEL_BYTES = policy["max_bytes"]

    job = db.get(MigrationJob, job_id)
    if not job or job.tenant_id != auth.tenant_id:
        raise HTTPException(status_code=404, detail="Migration job not found")
    name = (file.filename or "").lower()
    if not name.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail={"code": "UNSUPPORTED_TYPE", "allowed": [".xlsx", ".xlsm"]})
    raw = await file.read(MAX_EXCEL_BYTES + 1)
    if len(raw) > MAX_EXCEL_BYTES:
        raise HTTPException(status_code=413, detail={"code": "FILE_TOO_LARGE", "max_bytes": MAX_EXCEL_BYTES})
    if not raw.startswith(b"PK\x03\x04"):
        raise HTTPException(status_code=400, detail={"code": "INVALID_CONTENT", "message": "Not a valid xlsx (zip) file"})
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "Empty workbook")
    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    created = 0
    for row in rows[1:]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        data = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
        # Heuristic entity detection from headers
        if any(k in data for k in ("imo", "vessel", "vessel_name", "ship")):
            name = str(data.get("vessel_name") or data.get("vessel") or data.get("ship") or data.get("name") or "").strip()
            imo = str(data.get("imo") or "").strip() or None
            if name:
                db.add(
                    MigrationProposal(
                        job_id=job.id,
                        entity_type="vessel",
                        action="create",
                        confidence=0.9,
                        payload={"name": name, "imo": imo},
                        status="proposed",
                    )
                )
                created += 1
        elif any(k in data for k in ("counterparty", "charterer", "owner", "party")):
            name = str(
                data.get("counterparty")
                or data.get("charterer")
                or data.get("owner")
                or data.get("party")
                or data.get("name")
                or ""
            ).strip()
            if name:
                db.add(
                    MigrationProposal(
                        job_id=job.id,
                        entity_type="counterparty",
                        action="create",
                        confidence=0.86,
                        payload={"name": name, "type": str(data.get("type") or "other")},
                        status="proposed",
                    )
                )
                created += 1
        elif "unlocode" in data or "port" in data:
            name = str(data.get("port") or data.get("name") or "").strip()
            if name:
                db.add(
                    MigrationProposal(
                        job_id=job.id,
                        entity_type="port",
                        action="create",
                        confidence=0.9,
                        payload={
                            "name": name,
                            "unlocode": str(data.get("unlocode") or "") or None,
                            "timezone": str(data.get("timezone") or "UTC"),
                        },
                        status="proposed",
                    )
                )
                created += 1
    job.status = "review"
    job.summary = {**job.summary, "excel_proposals": created, "file": file.filename}
    db.commit()
    return {"status": job.status, "proposals_from_excel": created}


class MigrationCommitIn(BaseModel):
    proposal_ids: list[UUID]


@router.post("/settings/dataops/migrations/{job_id}/commit", tags=["DataOps"])
def commit_proposals(
    job_id: UUID,
    body: MigrationCommitIn,
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    from app.models_wave1 import Counterparty, Port, Vessel

    job = db.get(MigrationJob, job_id)
    if not job or job.tenant_id != auth.tenant_id:
        raise HTTPException(status_code=404, detail="Migration job not found")
    applied = {"vessel": 0, "counterparty": 0, "port": 0, "skipped": 0, "exists": 0}
    for pid in body.proposal_ids:
        prop = db.get(MigrationProposal, pid)
        if not prop or prop.job_id != job.id or prop.status != "proposed":
            applied["skipped"] += 1
            continue
        payload = prop.payload or {}
        if prop.entity_type == "vessel":
            imo = payload.get("imo")
            if imo and db.scalar(
                select(Vessel).where(Vessel.tenant_id == auth.tenant_id, Vessel.imo == str(imo)).limit(1)
            ):
                prop.status = "accepted"
                applied["exists"] += 1
                continue
            db.add(
                Vessel(
                    tenant_id=auth.tenant_id,
                    name=str(payload.get("name") or "Unknown"),
                    imo=str(imo) if imo else None,
                    vessel_type=payload.get("vessel_type"),
                    status="active",
                )
            )
            applied["vessel"] += 1
        elif prop.entity_type == "counterparty":
            name = str(payload.get("name") or "Unknown")
            if db.scalar(
                select(Counterparty).where(Counterparty.tenant_id == auth.tenant_id, Counterparty.name == name).limit(1)
            ):
                prop.status = "accepted"
                applied["exists"] += 1
                continue
            db.add(
                Counterparty(
                    tenant_id=auth.tenant_id,
                    name=name,
                    type=str(payload.get("type") or "other"),
                    country=payload.get("country"),
                    sanctions_status="clear",
                )
            )
            applied["counterparty"] += 1
        elif prop.entity_type == "port":
            unlocode = payload.get("unlocode")
            if unlocode and db.scalar(select(Port).where(Port.unlocode == str(unlocode)).limit(1)):
                prop.status = "accepted"
                applied["exists"] += 1
                continue
            db.add(
                Port(
                    name=str(payload.get("name") or "Unknown"),
                    unlocode=str(unlocode) if unlocode else None,
                    timezone=str(payload.get("timezone") or "UTC"),
                )
            )
            applied["port"] += 1
        else:
            applied["skipped"] += 1
            continue
        prop.status = "accepted"
    job.status = "committed"
    job.summary = {**job.summary, "commit": applied}
    db.commit()
    return {"status": "committed", "applied": applied}


@router.get(
    "/settings/dataops/migrations/{job_id}/proposals",
    response_model=list[MigrationProposalOut],
    tags=["DataOps"],
)
def list_proposals(
    job_id: UUID,
    auth: AuthContext = Depends(require_module("dataops")),
    db: Session = Depends(get_db),
):
    job = db.get(MigrationJob, job_id)
    if not job or job.tenant_id != auth.tenant_id:
        raise HTTPException(status_code=404, detail="Migration job not found")
    rows = db.scalars(select(MigrationProposal).where(MigrationProposal.job_id == job.id)).all()
    return [
        MigrationProposalOut(
            id=r.id,
            entity_type=r.entity_type,
            action=r.action,
            confidence=float(r.confidence),
            payload=r.payload,
            status=r.status,
        )
        for r in rows
    ]


@router.post("/settings/api-keys", response_model=ApiKeyOut, tags=["API Management"])
def create_api_key(
    body: ApiKeyCreate,
    auth: AuthContext = Depends(require_module("apim")),
    db: Session = Depends(get_db),
):
    raw = f"vos_{secrets.token_urlsafe(24)}"
    prefix = raw[:10]
    key_hash = hashlib.sha256(raw.encode()).hexdigest()
    row = ApiKey(
        tenant_id=auth.tenant_id,
        name=body.name,
        key_prefix=prefix,
        key_hash=key_hash,
        scopes=body.scopes,
        status="active",
        user_id=auth.user_id,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return ApiKeyOut(
        id=row.id,
        name=row.name,
        key_prefix=row.key_prefix,
        scopes=list(row.scopes or []),
        status=row.status,
        raw_key=raw,
    )


@router.get("/settings/api-keys", response_model=list[ApiKeyOut], tags=["API Management"])
def list_api_keys(
    auth: AuthContext = Depends(require_module("apim")),
    db: Session = Depends(get_db),
):
    rows = db.scalars(
        select(ApiKey).where(ApiKey.tenant_id == auth.tenant_id, ApiKey.status != "deleted")
    ).all()
    return [
        ApiKeyOut(
            id=r.id,
            name=r.name,
            key_prefix=r.key_prefix,
            scopes=list(r.scopes or []),
            status=r.status,
        )
        for r in rows
    ]


class ApiKeyUpdate(BaseModel):
    name: str | None = None
    scopes: list[str] | None = None
    status: str | None = None


@router.patch("/settings/api-keys/{key_id}", response_model=ApiKeyOut, tags=["API Management"])
def update_api_key(
    key_id: UUID,
    body: ApiKeyUpdate,
    auth: AuthContext = Depends(require_module("apim")),
    db: Session = Depends(get_db),
):
    row = db.get(ApiKey, key_id)
    if not row or row.tenant_id != auth.tenant_id or row.status == "deleted":
        raise HTTPException(404, "API key not found")
    if body.name is not None:
        row.name = body.name
    if body.scopes is not None:
        row.scopes = body.scopes
    if body.status is not None and body.status != "deleted":
        row.status = body.status
    db.commit()
    db.refresh(row)
    return ApiKeyOut(
        id=row.id,
        name=row.name,
        key_prefix=row.key_prefix,
        scopes=list(row.scopes or []),
        status=row.status,
    )


@router.delete("/settings/api-keys/{key_id}", tags=["API Management"])
def delete_api_key(
    key_id: UUID,
    auth: AuthContext = Depends(require_module("apim")),
    db: Session = Depends(get_db),
):
    row = db.get(ApiKey, key_id)
    if not row or row.tenant_id != auth.tenant_id or row.status == "deleted":
        raise HTTPException(404, "API key not found")
    soft_delete(
        db,
        tenant_id=auth.tenant_id,
        user_id=auth.user_id,
        entity_type="api_key",
        row=row,
        title=row.name,
    )
    db.commit()
    return {"ok": True, "recycled": True}
